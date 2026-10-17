"""Unit tests for multi-format document processing: Office docs, spreadsheets, standalone images, and HTML assets."""

import asyncio
import base64
import io
from pathlib import Path
from unittest.mock import MagicMock

import openpyxl
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from lib.doc_processor import (
    _format_pipe_table,
    convert_docx_to_markdown,
    convert_pptx_to_markdown,
    convert_csv_to_markdown,
    convert_xlsx_to_markdown,
    process_image_source,
    process_html_source,
)
from lib.utils import read_source_file, SUPPORTED_SOURCE_EXTENSIONS


def test_format_pipe_table():
    # Empty
    assert _format_pipe_table([]) == ""
    assert _format_pipe_table([[]]) == ""

    # Normal table
    rows = [
        ["Header 1", "Header 2"],
        ["Cell 1", "Cell 2"],
        ["Cell 3", "Cell 4 | with pipe\nand newline"],
    ]
    md = _format_pipe_table(rows)
    assert "| Header 1 | Header 2 |" in md
    assert "| :--- | :--- |" in md
    assert "| Cell 1 | Cell 2 |" in md
    assert "Cell 4 \\| with pipe and newline" in md

    # Truncation beyond max_rows
    large_rows = [["ID", "Value"]] + [[str(i), f"val_{i}"] for i in range(10)]
    md_truncated = _format_pipe_table(large_rows, max_rows=5)
    assert "[6 additional rows omitted for token limits]" in md_truncated


def test_convert_docx_to_markdown(tmp_path):
    docx_path = tmp_path / "research_notes.docx"
    doc = Document()
    doc.add_heading("Heading Level 1", level=1)
    doc.add_heading("Heading Level 2", level=2)
    doc.add_paragraph("This is a bullet point", style="List Bullet")
    doc.add_paragraph("Normal body paragraph explaining architecture.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Latency"
    table.cell(1, 0).text = "Ingest"
    table.cell(1, 1).text = "120ms"

    # Add an embedded image
    img = Image.new("RGB", (150, 150), color="blue")
    img_bytes = io.BytesIO()
    img.save(img_bytes, format="PNG")
    img_bytes.seek(0)
    doc.add_picture(img_bytes, width=Inches(2))

    doc.save(str(docx_path))

    assets_dir = tmp_path / "assets" / "research-notes"
    md = convert_docx_to_markdown(docx_path, assets_dir=assets_dir, doc_slug="research-notes")

    assert "# Research Notes" in md
    assert "## Heading Level 1" in md
    assert "### Heading Level 2" in md
    assert "- This is a bullet point" in md
    assert "Normal body paragraph explaining architecture." in md
    assert "| Component | Latency |" in md
    assert "| Ingest | 120ms |" in md
    assert "![Embedded Image" in md
    assert assets_dir.exists()
    assert len(list(assets_dir.glob("*.png"))) >= 1


def test_convert_pptx_to_markdown(tmp_path):
    pptx_path = tmp_path / "system_design.pptx"
    prs = Presentation()

    # Slide 1: Title and content
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "System Overview"
    slide1.placeholders[1].text = "Core design goals and architecture"

    # Slide 2: Table
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    table_shape = slide2.shapes.add_table(rows=2, cols=2, left=Inches(1), top=Inches(1), width=Inches(4), height=Inches(2))
    tbl = table_shape.table
    tbl.cell(0, 0).text = "Service"
    tbl.cell(0, 1).text = "Port"
    tbl.cell(1, 0).text = "MCP Router"
    tbl.cell(1, 1).text = "8765"

    prs.save(str(pptx_path))

    md = convert_pptx_to_markdown(pptx_path)
    assert "# Presentation: System Design" in md
    assert "## Slide 1: System Overview" in md
    assert "Core design goals and architecture" in md
    assert "## Slide 2" in md
    assert "| Service | Port |" in md
    assert "| MCP Router | 8765 |" in md


def test_convert_csv_and_tsv_to_markdown(tmp_path):
    # CSV
    csv_path = tmp_path / "benchmarks.csv"
    csv_path.write_text("Model,Param Count,MMLU\nClaude 3.5,Unknown,88.5\nGPT-4o,Unknown,88.7\n", encoding="utf-8")

    csv_md = convert_csv_to_markdown(csv_path)
    assert "# Data Table: Benchmarks" in csv_md
    assert "| Model | Param Count | MMLU |" in csv_md
    assert "| Claude 3.5 | Unknown | 88.5 |" in csv_md

    # TSV
    tsv_path = tmp_path / "data.tsv"
    tsv_path.write_text("Key\tValue\nHost\tlocalhost\nPort\t5432\n", encoding="utf-8")

    tsv_md = convert_csv_to_markdown(tsv_path)
    assert "# Data Table: Data" in tsv_md
    assert "| Key | Value |" in tsv_md
    assert "| Host | localhost |" in tsv_md


def test_convert_xlsx_to_markdown(tmp_path):
    xlsx_path = tmp_path / "company_metrics.xlsx"
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Financials"
    ws1.append(["Quarter", "Revenue", "EBITDA"])
    ws1.append(["Q1 2026", "$15M", "$3.5M"])
    ws1.append(["Q2 2026", "$18M", "$4.2M"])

    ws2 = wb.create_sheet(title="Headcount")
    ws2.append(["Department", "Count"])
    ws2.append(["Engineering", "45"])
    ws2.append(["Design", "8"])

    wb.save(str(xlsx_path))

    md = convert_xlsx_to_markdown(xlsx_path)
    assert "# Spreadsheet: Company Metrics" in md
    assert "## Sheet: Financials" in md
    assert "| Quarter | Revenue | EBITDA |" in md
    assert "| Q1 2026 | $15M | $3.5M |" in md
    assert "## Sheet: Headcount" in md
    assert "| Department | Count |" in md
    assert "| Engineering | 45 |" in md


def test_read_source_file_all_formats(tmp_path):
    # docx
    docx_path = tmp_path / "note.docx"
    doc = Document()
    doc.add_heading("Docx Note", level=1)
    doc.save(str(docx_path))
    assert "Docx Note" in read_source_file(docx_path)

    # csv
    csv_path = tmp_path / "table.csv"
    csv_path.write_text("A,B\n1,2\n", encoding="utf-8")
    assert "| A | B |" in read_source_file(csv_path)

    # image
    img_path = tmp_path / "diagram.png"
    img = Image.new("RGB", (50, 50), color="red")
    img.save(str(img_path))
    img_md = read_source_file(img_path)
    assert "# Image Asset: Diagram" in img_md
    assert "![Diagram](diagram.png)" in img_md


def test_process_image_source_with_vision(tmp_path):
    wiki_dir = tmp_path / "test-wiki"
    wiki_dir.mkdir()
    raw_dir = wiki_dir / "raw"
    raw_dir.mkdir()

    img_path = raw_dir / "network_topology.png"
    img = Image.new("RGB", (200, 200), color="green")
    img.save(str(img_path))

    mock_provider = MagicMock()
    mock_provider.call_vision.return_value = (
        "### Visual Analysis: network_topology.png\n"
        "- **Type:** Network Diagram\n"
        "- **Overview:** Gateway router connecting three internal subnets.\n"
    )

    md = asyncio.run(process_image_source(
        image_path=img_path,
        wiki_dir=wiki_dir,
        provider=mock_provider,
        domain_name="Computer Science",
    ))

    # Verification
    asset_file = wiki_dir / "wiki" / "assets" / "network-topology" / "network_topology.png"
    assert asset_file.exists()
    assert "# Network Topology" in md
    assert "![Network Topology](assets/network-topology/network_topology.png)" in md
    assert "> [!FIGURE: network_topology.png]" in md
    assert "Gateway router connecting three internal subnets." in md
    mock_provider.call_vision.assert_called_once()


def test_process_image_source_soft_fallback(tmp_path):
    wiki_dir = tmp_path / "test-wiki"
    wiki_dir.mkdir()
    img_path = wiki_dir / "raw" / "chart.jpg"
    img_path.parent.mkdir()
    img = Image.new("RGB", (200, 200), color="blue")
    img.save(str(img_path))

    mock_provider = MagicMock()
    mock_provider.call_vision.return_value = None  # Model quota exhausted

    md = asyncio.run(process_image_source(
        image_path=img_path,
        wiki_dir=wiki_dir,
        provider=mock_provider,
        domain_name="Finance",
    ))

    # Asset copied
    assert (wiki_dir / "wiki" / "assets" / "chart" / "chart.jpg").exists()
    assert "# Chart" in md
    assert "![Chart](assets/chart/chart.jpg)" in md
    assert "> [!FIGURE:" not in md


def test_process_html_source_with_base64_and_vision(tmp_path):
    wiki_dir = tmp_path / "test-wiki"
    wiki_dir.mkdir()
    raw_dir = wiki_dir / "raw"
    raw_dir.mkdir()

    # Create dummy base64 PNG (300x200 = 60000 pixels)
    img = Image.new("RGB", (300, 200), color="purple")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    b64_str = base64.b64encode(buf.getvalue()).decode("ascii")

    html_content = f"""
    <!DOCTYPE html>
    <html>
    <head><title>Obsidian Export</title></head>
    <body>
        <nav><a href="/home">Home</a></nav>
        <h1>Article With Diagrams</h1>
        <p>This is a paragraph about system design.</p>
        <img alt="System Flowchart" src="data:image/png;base64,{b64_str}">
        <div class="footer">Page footer</div>
    </body>
    </html>
    """
    html_path = raw_dir / "export.html"
    html_path.write_text(html_content, encoding="utf-8")

    mock_provider = MagicMock()
    mock_provider.call_vision.return_value = (
        "### Visual Analysis: embedded_img_1.png\n"
        "- **Type:** Architecture Flowchart\n"
        "- **Overview:** Shows pipeline flow from input to output."
    )

    md = asyncio.run(process_html_source(
        html_path=html_path,
        wiki_dir=wiki_dir,
        provider=mock_provider,
        domain_name="Architecture",
    ))

    # Asset extracted to disk
    assets_dir = wiki_dir / "wiki" / "assets" / "export"
    assert assets_dir.exists()
    extracted_imgs = list(assets_dir.glob("*.png"))
    assert len(extracted_imgs) == 1

    # Base64 string was replaced with file path reference
    assert "data:image/png;base64" not in md
    assert "assets/export/embedded_img_1.png" in md
    assert "Article With Diagrams" in md
    assert "> [!FIGURE: embedded_img_1.png]" in md
    assert "Shows pipeline flow from input to output." in md
